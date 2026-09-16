"""Intermediate pandas, Excel, and SQLite learning modules."""


def add_data_modules(section, unit):
    m=section('wrangling','Deeper pandas','Reshape, combine, and validate analytical tables.','Build a monthly regional report',[
        ('Long and wide tables','wide = df.pivot_table(index="month", columns="region", values="amount", aggfunc="sum", fill_value=0)\nlong = df.melt(id_vars="product", var_name="month", value_name="sales")','Choose a layout for the next calculation. pivot_table aggregates duplicate combinations; pivot requires each combination to be unique.'),
        ('Validate joins','combined = orders.merge(lookup, on="id", how="left", validate="many_to_one", indicator=True)','Check key cardinality and unmatched records before trusting the combined table.'),
        ('Time series','daily = df.set_index("date")["amount"].resample("D").sum()','Parse dates and sort them before time-based operations. Distinguish absent dates from recorded zero values.')])
    unit(m,'pivot','Reshape records into a comparison table',
        'Long-form records store one observation per row. A pivot table spreads one category across columns, making comparisons easier. Repeated month-region combinations need an aggregation rule; sum combines their amounts. fill_value=0 applies only when an absent combination should mean no recorded sales under this exercise’s policy.',
        'wide = df.pivot_table(index="month", columns="region", values="amount", aggfunc="sum", fill_value=0)',
        ['Create `wide` indexed by `month` with `North` and `South` columns.', 'Sum duplicate combinations and fill absent combinations with zero. Sort months.'],
        'import pandas as pd\ndf=pd.DataFrame({"month":["Jan","Jan","Jan","Feb"],"region":["North","North","South","North"],"amount":[10,5,20,30]})',
        'import pandas as pd\ndf=pd.DataFrame({"month":["Jan","Jan","Jan","Feb"],"region":["North","North","South","North"],"amount":[10,5,20,30]})\nwide=df.pivot_table(index="month",columns="region",values="amount",aggfunc="sum",fill_value=0).sort_index()',
        [('Aggregate duplicate keys','assert wide.loc["Jan","North"]==15 and wide.loc["Jan","South"]==20'),('Fill absent combinations','assert wide.loc["Feb","South"]==0 and wide.loc["Feb","North"]==30')],
        ['Use pivot_table rather than pivot.','Choose aggfunc="sum".','Set fill_value=0.'],
        'The comparison table contains the same revenue in a different layout. Alphabetical month names are not chronological; use date labels for time ordering.',review=['pandas-aggregations'])
    unit(m,'melt','Turn month columns into records',
        'A wide spreadsheet often uses a separate column for each month. That layout is convenient for reading but awkward for grouping many periods. melt keeps identifier columns and turns the remaining column names into values. Specify the identifier and output names so the resulting schema is predictable.',
        'long = df.melt(id_vars="product", var_name="month", value_name="sales")',
        ['Create `long` with `product`, `month`, and `sales` columns.', 'Sort by `product` then `month`, and reset its index.'],
        'import pandas as pd\ndf=pd.DataFrame({"product":["Tea","Cake"],"Jan":[10,20],"Feb":[15,25]})',
        'import pandas as pd\ndf=pd.DataFrame({"product":["Tea","Cake"],"Jan":[10,20],"Feb":[15,25]})\nlong=df.melt(id_vars="product",var_name="month",value_name="sales").sort_values(["product","month"]).reset_index(drop=True)',
        [('Produce four observations','assert len(long)==4 and long.columns.tolist()==["product","month","sales"]'),('Retain every measurement','assert long.to_dict("records")==[{"product":"Cake","month":"Feb","sales":25},{"product":"Cake","month":"Jan","sales":20},{"product":"Tea","month":"Feb","sales":15},{"product":"Tea","month":"Jan","sales":10}]')],
        ['Keep product as the identifier.','Name the new columns month and sales.','Sort after reshaping.'],
        'Each sales observation is now a row, making future periods easier to append.',review=['pandas-rename'])
    unit(m,'join-audit','Expose unmatched lookup keys',
        'A join can silently drop records or multiply them. A left merge preserves orders, validate checks the expected key relationship, and indicator records whether a match was found. Audit unmatched keys before calculating a summary. A missing lookup match is not automatically a zero price.',
        'joined=orders.merge(products,on="id",how="left",validate="many_to_one",indicator=True)',
        ['Write `join_orders(orders, products)`, returning `(joined, missing_ids)`.', 'Use a left merge validated as many-to-one. Return unique sorted unmatched IDs. Duplicate product keys must raise `pd.errors.MergeError`.'],
        'import pandas as pd\n\ndef join_orders(orders, products):\n    pass',
        'import pandas as pd\n\ndef join_orders(orders,products):\n    joined=orders.merge(products,on="id",how="left",validate="many_to_one",indicator=True)\n    missing_ids=sorted(joined.loc[joined["_merge"]=="left_only","id"].unique().tolist())\n    return joined,missing_ids',
        [('Preserve unmatched orders','orders=pd.DataFrame({"id":[1,9,9],"qty":[2,1,3]}); products=pd.DataFrame({"id":[1],"price":[5]})\nj,missing=join_orders(orders,products)\nassert len(j)==3 and missing==[9]'),('Reject duplicate lookup keys','try:\n    join_orders(orders,pd.DataFrame({"id":[1,1],"price":[5,6]}))\nexcept pd.errors.MergeError:\n    pass\nelse:\n    raise AssertionError()')],
        ['Set validate="many_to_one".','indicator=True adds _merge.','Filter left_only keys and deduplicate them.'],
        'A validated join protects both row counts and the interpretation of missing data.',review=['pandas-merge','collections-sets'])
    unit(m,'concat','Combine exports and remove repeated records',
        'concat stacks compatible tables. It does not automatically remove duplicated records. When exports overlap, first define a business key and a rule for which record to keep. Here the second export is newer, so the final occurrence of each order_id wins. Reset the resulting index after deduplication.',
        'combined=pd.concat([older,newer],ignore_index=True)',
        ['Write `combine_exports(older,newer)`, retaining the last occurrence of each `order_id`.', 'Return rows sorted by `order_id` with a fresh index. Do not modify the inputs.'],
        'import pandas as pd\n\ndef combine_exports(older,newer):\n    pass',
        'import pandas as pd\n\ndef combine_exports(older,newer):\n    return pd.concat([older,newer],ignore_index=True).drop_duplicates("order_id",keep="last").sort_values("order_id").reset_index(drop=True)',
        [('Retain the newer record','a=pd.DataFrame({"order_id":[1,2],"amount":[10,20]}); b=pd.DataFrame({"order_id":[2,3],"amount":[25,30]})\nr=combine_exports(a,b)\nassert r.to_dict("records")==[{"order_id":1,"amount":10},{"order_id":2,"amount":25},{"order_id":3,"amount":30}]'),('Preserve input and index','assert a["amount"].tolist()==[10,20] and r.index.tolist()==[0,1,2]')],
        ['Concatenate older before newer.','Drop duplicates on order_id with keep="last".','Sort then reset the index.'],
        'Combining files and resolving overlapping records are distinct steps.',review=['pandas-deduplicate','files-merge-review'])
    unit(m,'rolling','Calculate a chronological rolling average',
        'Time-based analysis depends on ordering. Parse the date column, sort records, then apply a rolling window. A three-observation window is different from a three-calendar-day window when dates are missing. This exercise explicitly uses three observations and requires all three before returning a mean.',
        'df["average"]=df["amount"].rolling(3,min_periods=3).mean()',
        ['Create `ordered` sorted by parsed `date`.', 'Add `average` as the three-observation rolling mean of `amount`. Keep the first two means missing.'],
        'import pandas as pd\ndf=pd.DataFrame({"date":["2026-01-03","2026-01-01","2026-01-04","2026-01-02"],"amount":[30,10,40,20]})',
        'import pandas as pd\ndf=pd.DataFrame({"date":["2026-01-03","2026-01-01","2026-01-04","2026-01-02"],"amount":[30,10,40,20]})\nordered=df.copy()\nordered["date"]=pd.to_datetime(ordered["date"])\nordered=ordered.sort_values("date").reset_index(drop=True)\nordered["average"]=ordered["amount"].rolling(3,min_periods=3).mean()',
        [('Sort before rolling','assert ordered["amount"].tolist()==[10,20,30,40]'),('Require complete windows','assert ordered["average"].iloc[:2].isna().all() and ordered["average"].iloc[2:].tolist()==[20,30]')],
        ['Parse dates before sorting.','Use rolling(3,min_periods=3).','Do not replace initial missing means with zero.'],
        'The initial missing results mean a complete window is not yet available.',review=['analysis-rolling','pandas-monthly-review'])
    unit(m,'monthly-project','Project · clean, group, and explain exclusions',
        'Build a monthly regional report from a deliberately imperfect CSV export. The project combines date parsing, numeric coercion, text cleaning, and grouping. Reject invalid dates, invalid or negative amounts, and blank regions. Keep valid zero amounts. Export both a tidy report and a small quality summary so the reader knows what was excluded.',
        'summary.to_csv("monthly.csv",index=False)',
        ['Read `transactions.csv`. Clean `region` to lowercase and parse `date` and `amount`.', 'Create `summary` with `month`, `region`, `total` sorted by `month` and `region`.', 'Write `monthly.csv` and `quality.json` containing `accepted` and `rejected` counts.'],
        'import pandas as pd\nimport json',
        '''
        import pandas as pd
        import json
        df=pd.read_csv('transactions.csv',keep_default_na=False)
        df['date']=pd.to_datetime(df['date'],errors='coerce')
        df['amount']=pd.to_numeric(df['amount'],errors='coerce')
        df['region']=df['region'].str.strip().str.lower()
        valid=df['date'].notna() & df['amount'].notna() & (df['amount']>=0) & (df['region']!='')
        quality={'accepted':int(valid.sum()),'rejected':int((~valid).sum())}
        clean=df.loc[valid].copy()
        clean['month']=clean['date'].dt.strftime('%Y-%m')
        summary=clean.groupby(['month','region']).agg(total=('amount','sum')).reset_index().sort_values(['month','region']).reset_index(drop=True)
        summary.to_csv('monthly.csv',index=False)
        with open('quality.json','w',encoding='utf-8') as file:
            json.dump(quality,file)
        ''',
        [('Report correct groups','assert summary.to_dict("records")==[{"month":"2026-01","region":"north","total":30},{"month":"2026-02","region":"south","total":0}]'),('Audit exclusions','assert quality=={"accepted":3,"rejected":3}'),('Export both results','assert pd.read_csv("monthly.csv").to_dict("records")==summary.to_dict("records")\nassert json.load(open("quality.json"))==quality')],
        ['Construct one Boolean validity mask.','Count its True and False values before filtering.','Group cleaned data by both month and region.'],
        'You now have a reproducible transformation and an explicit data-quality report.',kind='project',minutes=35,review=['bridge-quality'],files={'transactions.csv':'date,region,amount\n2026-01-01, North ,10\n2026-01-03,NORTH,20\n2026-02-01,South,0\nbad,North,999\n2026-02-02,South,-4\n2026-02-03, ,8\n'},exports=['monthly.csv','quality.json'])

    m=section('excel','Excel report automation','Read, validate, and create real workbooks.','Create a monthly workbook',[
        ('Read and write','wb=load_workbook("source.xlsx",data_only=False)\nws=wb["Sales"]\nwb.save("report.xlsx")','openpyxl works with .xlsx files locally. data_only=False reads formulas; data_only=True reads cached results when available.'),
        ('Rows and formatting','ws.append(["Product","Revenue"])\nws["B2"].number_format="0.00"','Use values_only=True when you need values rather than cell objects. Formatting changes appearance, not stored numeric values.'),
        ('Formula limits','ws["C2"]="=A2*B2"','openpyxl writes formulas but does not calculate them. Excel calculates formulas when opened; use explicit Python totals for immediately inspectable exports.')])
    setup='''
from openpyxl import Workbook
wb=Workbook(); ws=wb.active; ws.title='Sales'
ws.append(['Product','Qty','Price'])
ws.append(['Tea',2,5]); ws.append(['Cake',3,8])
wb.save('source.xlsx')
'''
    unit(m,'read','Read rows from a workbook',
        'A workbook contains named worksheets. Select the worksheet explicitly rather than assuming the active sheet is the one you want. iter_rows can return plain values and skip the header with min_row=2. The fixture creates a real workbook before your code runs. You can inspect source.xlsx by reading its rows and printing them.',
        'from openpyxl import load_workbook\nwb=load_workbook("source.xlsx")\nrows=list(wb["Sales"].iter_rows(min_row=2,values_only=True))',
        ['Read `source.xlsx`, worksheet `Sales`, into `rows` without the header.', 'Calculate `revenue` as `qty * price` across the `rows`. Close the workbook.'],
        'from openpyxl import load_workbook',
        'from openpyxl import load_workbook\nwb=load_workbook("source.xlsx")\nrows=list(wb["Sales"].iter_rows(min_row=2,values_only=True))\nrevenue=sum(qty*price for name,qty,price in rows)\nwb.close()',
        [('Read the intended sheet','assert rows==[("Tea",2,5),("Cake",3,8)]'),('Calculate from numeric cells','assert revenue==34')],
        ['Select wb["Sales"].','Skip row 1 and ask for values_only=True.','Unpack name, qty, price for the calculation.'],
        'The workbook’s numeric cells arrive as numbers, so conversion is not needed for this clean fixture.',setup=setup,review=['files-csv'])
    unit(m,'write','Build a workbook from records',
        'A generated workbook should have a clear sheet name and header row. append adds a row of values. Keep numbers numeric so the recipient can calculate with them later. Save the workbook to a declared output name to make it downloadable from Codey. The checker opens the generated workbook and inspects its actual cells.',
        'wb=Workbook()\nws=wb.active\nws.title="Summary"\nws.append(["Product","Revenue"])',
        ['Create `report.xlsx` with a `Summary` sheet.', 'Write `Product`, `Revenue` headers followed by `Tea,30` and `Cake,24`.'],
        'from openpyxl import Workbook',
        'from openpyxl import Workbook\nwb=Workbook(); ws=wb.active; ws.title="Summary"\nws.append(["Product","Revenue"]); ws.append(["Tea",30]); ws.append(["Cake",24])\nwb.save("report.xlsx")',
        [('Write the expected cells','from openpyxl import load_workbook\ncheck=load_workbook("report.xlsx"); assert list(check["Summary"].values)==[("Product","Revenue"),("Tea",30),("Cake",24)]; check.close()')],
        ['Rename the active sheet.','Append the header and both data rows.','Save with wb.save("report.xlsx").'],
        'The result is a real spreadsheet that can be downloaded and opened outside the lesson.',exports=['report.xlsx'])
    unit(m,'format','Make a workbook readable',
        'Formatting supports interpretation without changing values. A bold header, sensible column widths, and a frozen header row make an export easier to use. Number formats determine how decimals are displayed. Apply formatting to cells and worksheet settings, then save; edits to an in-memory workbook are not persisted until it is saved.',
        'from openpyxl.styles import Font\nws["A1"].font=Font(bold=True)\nws.freeze_panes="A2"',
        ['Open `source.xlsx` and bold all three header cells.', 'Freeze at `A2`, set column `A` width to 22, and format `Price` cells as `0.00`.', 'Save `formatted.xlsx`.'],
        'from openpyxl import load_workbook\nfrom openpyxl.styles import Font',
        'from openpyxl import load_workbook\nfrom openpyxl.styles import Font\nwb=load_workbook("source.xlsx"); ws=wb["Sales"]\nfor cell in ws[1]:\n    cell.font=Font(bold=True)\nws.freeze_panes="A2"; ws.column_dimensions["A"].width=22\nfor row in ws.iter_rows(min_row=2,min_col=3,max_col=3):\n    row[0].number_format="0.00"\nwb.save("formatted.xlsx"); wb.close()',
        [('Persist worksheet usability settings','from openpyxl import load_workbook\nw=load_workbook("formatted.xlsx"); s=w["Sales"]\nassert all(c.font.bold for c in s[1]) and s.freeze_panes=="A2" and s.column_dimensions["A"].width==22'),('Format prices without changing numbers','assert s["C2"].value==5 and s["C2"].number_format=="0.00" and s["C3"].number_format=="0.00"\nw.close()')],
        ['Loop through ws[1] for the header.','freeze_panes selects the first unfrozen cell.','Set number_format on the price cells before saving.'],
        'The underlying values remain numeric, while the workbook is easier to scan.',setup=setup,exports=['formatted.xlsx'])
    unit(m,'formulas','Write formulas and know their limits',
        'Excel formulas are strings beginning with =. openpyxl stores them but does not evaluate them. Include a Python-calculated total when immediate verification is important, and let Excel recalculate the formulas when the workbook is opened. Reading with data_only=True may return None if no cached calculation exists yet.',
        'ws["D2"]="=B2*C2"\npython_total=2*5',
        ['Add a `Revenue` header in `D1` and formulas `=B2*C2` and `=B3*C3`.', 'Calculate `python_total` from the original numeric cells.', 'Save `formulas.xlsx`.'],
        'from openpyxl import load_workbook',
        'from openpyxl import load_workbook\nwb=load_workbook("source.xlsx"); ws=wb["Sales"]\nws["D1"]="Revenue"\nws["D2"]="=B2*C2"; ws["D3"]="=B3*C3"\npython_total=sum(ws.cell(r,2).value*ws.cell(r,3).value for r in [2,3])\nwb.save("formulas.xlsx"); wb.close()',
        [('Store genuine formula cells','from openpyxl import load_workbook\nw=load_workbook("formulas.xlsx",data_only=False); s=w["Sales"]\nassert s["D1"].value=="Revenue" and s["D2"].value=="=B2*C2" and s["D3"].value=="=B3*C3"\nw.close()'),('Compute an independent total','assert python_total==34')],
        ['Assign formula text beginning with =.','Use the original cells for python_total.','The formula text and Python total are checked separately.'],
        'Writing a formula is not the same as evaluating it. The report makes that distinction explicit.',setup=setup,exports=['formulas.xlsx'])
    unit(m,'validate','Separate valid rows from rejected ones',
        'Real spreadsheets can contain blanks and text in numeric columns. Validate rows before calculation and keep a rejection count. This fixture expects a nonblank product and nonnegative numeric quantity and price. A blank cell is None; a numeric zero is valid. Do not silently coerce an unknown quantity into zero.',
        'valid = isinstance(qty,(int,float)) and qty >= 0',
        ['Read `Dirty` from `source.xlsx`.', 'Create `valid_rows` of accepted `(product, qty, price)` tuples and `rejected` as a count.', 'Strip `product` names and preserve input order.'],
        'from openpyxl import load_workbook',
        '''
        from openpyxl import load_workbook
        wb=load_workbook('source.xlsx'); valid_rows=[]; rejected=0
        for product,qty,price in wb['Dirty'].iter_rows(min_row=2,values_only=True):
            if isinstance(product,str) and product.strip() and isinstance(qty,(int,float)) and qty>=0 and isinstance(price,(int,float)) and price>=0:
                valid_rows.append((product.strip(),qty,price))
            else:
                rejected+=1
        wb.close()
        ''',
        [('Keep valid zero and ordinary rows','assert valid_rows==[("Tea",2,5),("Cake",0,8)]'),('Count incomplete and invalid rows','assert rejected==3')],
        ['Inspect each cell type before calculating.','Use explicit >=0 checks so zero survives.','Count every row that fails the contract.'],
        'The cleaned records are ready for reporting, and exclusions remain visible.',setup="from openpyxl import Workbook\nwb=Workbook(); ws=wb.active; ws.title='Dirty'\nfor row in [('Product','Qty','Price'),(' Tea ',2,5),('Cake',0,8),('Coffee','many',4),(None,2,3),('Juice',-1,5)]:\n    ws.append(row)\nwb.save('source.xlsx')",review=['bridge-refunds'])
    unit(m,'workbook-project','Project · a report with an audit sheet',
        'Produce a workbook someone can use without reading your Python code. One sheet should contain the cleaned sales summary; another should document how many rows were accepted or rejected. This small project combines CSV input, validation, grouping, and workbook formatting. Use fixed numeric outputs so the report is immediately inspectable.',
        'audit=wb.create_sheet("Audit")\naudit.append(["Metric","Count"])',
        ['Read `sales.csv`. Reject invalid or negative `qty`, and blank `product` names. Strip `product` names.', 'Create `Summary` with `Product,Units` headers and alphabetically sorted product totals.', 'Create `Audit` with `Metric,Count` headers, then `accepted` and `rejected` rows. Bold both headers and save `sales_report.xlsx`.'],
        'import csv\nfrom openpyxl import Workbook\nfrom openpyxl.styles import Font',
        '''
        import csv
        from openpyxl import Workbook
        from openpyxl.styles import Font
        totals={}; accepted=0; rejected=0
        with open('sales.csv',newline='',encoding='utf-8') as file:
            for row in csv.DictReader(file):
                try:
                    product=row['product'].strip(); qty=int(row['qty'])
                    if not product or qty<0: raise ValueError()
                except ValueError:
                    rejected+=1; continue
                totals[product]=totals.get(product,0)+qty; accepted+=1
        wb=Workbook(); ws=wb.active; ws.title='Summary'; ws.append(['Product','Units'])
        for name,units in sorted(totals.items()): ws.append([name,units])
        audit=wb.create_sheet('Audit'); audit.append(['Metric','Count']); audit.append(['accepted',accepted]); audit.append(['rejected',rejected])
        for sheet in [ws,audit]:
            for cell in sheet[1]: cell.font=Font(bold=True)
            sheet.freeze_panes='A2'
        wb.save('sales_report.xlsx')
        ''',
        [('Create correct summary rows','from openpyxl import load_workbook\nw=load_workbook("sales_report.xlsx"); assert list(w["Summary"].values)==[("Product","Units"),("Cake",0),("Tea",5)]'),('Include audit and header formatting','assert list(w["Audit"].values)==[("Metric","Count"),("accepted",3),("rejected",2)]\nassert all(c.font.bold for s in w for c in s[1]); w.close()')],
        ['Validate each source row before accumulating.','Build two sheets with explicit headers.','Write Python-calculated totals and counts, then save.'],
        'The report contains both useful output and enough context to understand its limitations.',kind='project',minutes=35,files={'sales.csv':'product,qty\nTea,2\n Tea ,3\nCake,0\nCoffee,nope\nJuice,-1\n'},exports=['sales_report.xlsx'],review=['files-merge-review','excel-format'])

    m=section('sql','SQL with SQLite','Query and update a local relational database.','Build a customer revenue report',[
        ('Practice database schema', 'customers(id INTEGER PRIMARY KEY, name TEXT, region TEXT)\norders(id INTEGER PRIMARY KEY, customer_id INTEGER, amount REAL)', 'shop.db is already created for every attempt. customers contains Ari (1, North), Bo (2, South), and Cy (3, North). orders contains (id, customer_id, amount): (1,1,10), (2,1,20), (3,2,40). orders.customer_id refers to customers.id.'),
        ('Query and bind values','rows=conn.execute("SELECT name FROM customers WHERE region = ?", (region,)).fetchall()','Use ? placeholders for values. Never interpolate user input into SQL. SQLite runs locally and requires no account or service.'),
        ('Group and join','SELECT customer_id, SUM(amount) FROM orders GROUP BY customer_id\nSELECT c.name,o.amount FROM customers c JOIN orders o ON c.id=o.customer_id','A join connects related keys. Aggregation reduces rows into groups. ORDER BY makes result order explicit.'),
        ('Transactions','with conn:\n    conn.execute("INSERT INTO items(name) VALUES (?)", (name,))','The connection context commits on success and rolls back on an exception. It does not close the connection; close it separately.')])
    setup="import sqlite3\nc=sqlite3.connect('shop.db')\nc.executescript('CREATE TABLE customers(id INTEGER PRIMARY KEY,name TEXT,region TEXT); CREATE TABLE orders(id INTEGER PRIMARY KEY,customer_id INTEGER,amount REAL);')\nc.executemany('INSERT INTO customers VALUES (?,?,?)',[(1,'Ari','North'),(2,'Bo','South'),(3,'Cy','North')])\nc.executemany('INSERT INTO orders VALUES (?,?,?)',[(1,1,10),(2,1,20),(3,2,40)])\nc.commit(); c.close()"
    unit(m,'select','Ask a database for selected rows',
        'SQL describes the rows you want rather than the loop used to find them. SELECT chooses columns, WHERE filters records, and ORDER BY defines the result order. sqlite3 returns rows as tuples by default. Always specify ordering when the order matters; a database table has no guaranteed presentation order.',
        'rows=conn.execute("SELECT name FROM customers ORDER BY name").fetchall()',
        ['Connect to `shop.db` and select names of `North` customers, alphabetically.', 'Store the returned tuples in `rows` and close the connection.'],
        'import sqlite3',
        'import sqlite3\nconn=sqlite3.connect("shop.db")\nrows=conn.execute("SELECT name FROM customers WHERE region = ? ORDER BY name",("North",)).fetchall()\nconn.close()',
        [('Filter and order database records','assert rows==[("Ari",),("Cy",)]')],
        ['Connect with sqlite3.connect.','Use WHERE region = ? with a one-item tuple.','Call fetchall before closing.'],
        'Filtering happened in the database and returned only the requested column.',setup=setup,review=['pandas-select'])
    unit(m,'parameters','Bind an input value safely',
        'Parameter binding keeps data separate from SQL syntax. A customer name may contain an apostrophe; that should remain ordinary text, not break the query. A one-item parameter tuple needs a trailing comma. This lesson tests both a real name and SQL-like text to ensure the input is treated as a value.',
        'conn.execute("SELECT id FROM customers WHERE name = ?", (name,))',
        ['Write `find_customer(conn,name)`, returning matching `(id, name)` tuples sorted by `id`.', 'Use a parameterized query. A SQL-like name should return no records.'],
        'import sqlite3\n\ndef find_customer(conn,name):\n    pass',
        'import sqlite3\n\ndef find_customer(conn,name):\n    return conn.execute("SELECT id,name FROM customers WHERE name = ? ORDER BY id",(name,)).fetchall()',
        [('Find an ordinary name',"c=sqlite3.connect('shop.db'); assert find_customer(c,'Ari')==[(1,'Ari')]"),('Handle special characters as data',"c.execute('INSERT INTO customers VALUES (?,?,?)',(4,\"O'Neil\",'West')); assert find_customer(c,\"O'Neil\")==[(4,\"O'Neil\")]\nassert find_customer(c,\"' OR 1=1 --\")==[]; c.close()")],
        ['Keep ? in the SQL string.','Pass (name,) separately.','Do not build SQL with an f-string.'],
        'Binding handles quoting and prevents the input from changing query structure.',setup=setup,review=['functions-return'])
    unit(m,'group','Aggregate inside SQL',
        'GROUP BY combines records that share a key. SUM calculates a total within each group, while HAVING filters groups after aggregation. WHERE instead filters individual rows before grouping. Keep those stages separate in your reasoning. Here each customer’s orders are grouped and only sufficiently large totals remain.',
        'SELECT customer_id, SUM(amount) AS total FROM orders GROUP BY customer_id HAVING SUM(amount) >= 30',
        ['Create `totals` as `(customer_id, total)` tuples for totals of at least 30.', 'Sort by `total` descending, then `customer_id` ascending.'],
        'import sqlite3\nconn=sqlite3.connect("shop.db")',
        'import sqlite3\nconn=sqlite3.connect("shop.db")\ntotals=conn.execute("SELECT customer_id,SUM(amount) AS total FROM orders GROUP BY customer_id HAVING SUM(amount)>=30 ORDER BY total DESC,customer_id").fetchall()\nconn.close()',
        [('Group, filter, and rank','assert totals==[(2,40),(1,30)]')],
        ['Group by customer_id.','Filter the aggregate with HAVING.','ORDER BY total DESC gives the ranking.'],
        'SQL can return an analytical summary without transferring every input row into Python.',setup=setup,review=['analysis-group-means'])
    unit(m,'left-join','Keep customers without orders',
        'A left join preserves customers even when no matching order exists. SUM over the missing side returns NULL, so COALESCE supplies zero under the stated policy that no orders means zero revenue. COUNT(order_id) counts matched orders; COUNT(*) would also count the preserved customer row.',
        'SELECT c.name, COALESCE(SUM(o.amount),0) FROM customers c LEFT JOIN orders o ON c.id=o.customer_id GROUP BY c.id',
        ['Create `report` tuples of `name`, `total`, `order_count` for all customers.', 'Use zero total and zero count for a customer without orders. Sort by `name`.'],
        'import sqlite3\nconn=sqlite3.connect("shop.db")',
        'import sqlite3\nconn=sqlite3.connect("shop.db")\nreport=conn.execute("SELECT c.name,COALESCE(SUM(o.amount),0),COUNT(o.id) FROM customers c LEFT JOIN orders o ON c.id=o.customer_id GROUP BY c.id,c.name ORDER BY c.name").fetchall()\nconn.close()',
        [('Preserve unmatched customers','assert report==[("Ari",30,2),("Bo",40,1),("Cy",0,0)]')],
        ['Start from customers and LEFT JOIN orders.','Use COALESCE(SUM(...),0).','Count o.id instead of *.'],
        'The report distinguishes no orders from a customer missing from the report entirely.',setup=setup,review=['wrangling-join-audit'])
    unit(m,'transaction','Make a batch update atomic',
        'A transaction makes a set of writes succeed or fail together. If an invalid value is found halfway through a batch, earlier writes in that transaction should be rolled back. Raising the error inside the connection context triggers rollback. Catching it and pretending the operation succeeded would leave callers with a false picture of the data.',
        'with conn:\n    conn.execute("INSERT INTO orders VALUES (?,?,?)", row)',
        ['Write `add_orders(conn,rows)`, inserting `(id, customer_id, amount)` records.', 'Reject negative amounts with `ValueError`. The whole batch must roll back if any row is invalid.'],
        'import sqlite3\n\ndef add_orders(conn,rows):\n    pass',
        'import sqlite3\n\ndef add_orders(conn,rows):\n    with conn:\n        for row in rows:\n            if row[2]<0:\n                raise ValueError("amount must be nonnegative")\n            conn.execute("INSERT INTO orders VALUES (?,?,?)",row)',
        [('Commit valid rows','c=sqlite3.connect("shop.db"); add_orders(c,[(4,3,5)]); assert c.execute("SELECT amount FROM orders WHERE id=4").fetchone()==(5,)'),('Roll back the whole invalid batch','try:\n    add_orders(c,[(5,3,7),(6,3,-1)])\nexcept ValueError:\n    pass\nelse:\n    raise AssertionError()\nassert c.execute("SELECT COUNT(*) FROM orders WHERE id IN (5,6)").fetchone()[0]==0; c.close()')],
        ['Place the entire loop inside with conn:.','Raise before inserting a negative amount.','Let the exception leave the with block.'],
        'Atomic updates prevent partially applied batches from looking complete.',setup=setup,review=['functions-raise'])
    unit(m,'sql-project','Project · export a customer report',
        'Combine a relational query with a portable report. Include every customer, calculate their order total and count, and rank by revenue with a stable alphabetical tie-break. Write one CSV header and the query results. This project has no starting query: use the relationship between customers and orders to build it.',
        'writer.writerow(["name","region","revenue","orders"])\nwriter.writerows(rows)',
        ['Write `customer_report.csv` with `name,region,revenue,orders` columns.', 'Include zero-order customers. Rank by `revenue` descending, then `name` ascending.', 'Set `grand_total` to the sum of report `revenue`.'],
        'import sqlite3\nimport csv',
        '''
        import sqlite3
        import csv
        conn=sqlite3.connect('shop.db')
        rows=conn.execute('SELECT c.name,c.region,COALESCE(SUM(o.amount),0) AS revenue,COUNT(o.id) FROM customers c LEFT JOIN orders o ON c.id=o.customer_id GROUP BY c.id,c.name,c.region ORDER BY revenue DESC,c.name').fetchall()
        grand_total=sum(row[2] for row in rows)
        with open('customer_report.csv','w',newline='',encoding='utf-8') as file:
            writer=csv.writer(file); writer.writerow(['name','region','revenue','orders']); writer.writerows(rows)
        conn.close()
        ''',
        [('Rank all customers','assert [(r[0],r[2],r[3]) for r in rows]==[("Bo",40,1),("Ari",30,2),("Cy",0,0)] and grand_total==70'),('Export the report schema','with open("customer_report.csv",newline="",encoding="utf-8") as f:\n    reader=csv.DictReader(f); assert reader.fieldnames==["name","region","revenue","orders"]; data=list(reader)\nassert [r["name"] for r in data]==["Bo","Ari","Cy"] and sum(float(r["revenue"]) for r in data)==70')],
        ['Use the left-join pattern with grouping.','Alias the total to order by it.','Write the returned tuples with csv.writer.'],
        'The exported report connects database querying with the file workflows you already know.',kind='project',minutes=35,setup=setup,exports=['customer_report.csv'],review=['files-write-csv','sql-left-join'])
